#!/usr/bin/env python3
"""Convert New Test list updated.xlsx -> database/seeders/data/lab_tests.json"""

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "New Test list  updated.xlsx"
OUT = ROOT / "database" / "seeders" / "data" / "lab_tests.json"


def normalize_head(value) -> str:
    if value is None:
        return "General Pathology"
    head = str(value).strip()
    if head == "" or head.lower() in ("none", "carry_out", "null"):
        return "General Pathology"
    return head


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    tests = []
    seen_ids = set()

    for row in rows:
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue
        try:
            test_id = int(float(row[0]))
        except (TypeError, ValueError):
            continue

        name = str(row[1] or "").strip()
        if name == "":
            continue

        if test_id in seen_ids:
            continue
        seen_ids.add(test_id)

        price = row[2]
        try:
            price = float(price) if price is not None else 0.0
        except (TypeError, ValueError):
            price = 0.0

        test_type = str(row[3] or "Routine").strip() or "Routine"
        head = normalize_head(row[4])
        report = str(row[5] or "Same Day").strip() or "Same Day"

        tests.append(
            {
                "id": test_id,
                "name": name,
                "price": price,
                "type": test_type,
                "head": head,
                "report": report,
            }
        )

    tests.sort(key=lambda t: t["id"])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(tests, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Wrote {len(tests)} tests -> {OUT}")


if __name__ == "__main__":
    main()
